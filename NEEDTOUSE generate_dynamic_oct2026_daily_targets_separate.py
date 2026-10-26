import os
import sys
import time
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paths
WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(WORKSPACE_DIR, "12_10_2026")
MONTHLY_TARGET_FILE = os.path.join(DATA_DIR, "branch_wise_monthly_target_sample.csv")

OUTPUT_EXCEL_FILENAME = "branch_wise_daily_target_oct2026_dynamic_separate.xlsx"
OUTPUT_CSV_FILENAME = "branch_wise_daily_target_oct2026_dynamic_separate.csv"

OUTPUT_EXCEL_PATH_1 = os.path.join(DATA_DIR, OUTPUT_EXCEL_FILENAME)
OUTPUT_EXCEL_PATH_2 = os.path.join(WORKSPACE_DIR, OUTPUT_EXCEL_FILENAME)

OUTPUT_CSV_PATH_1 = os.path.join(DATA_DIR, OUTPUT_CSV_FILENAME)
OUTPUT_CSV_PATH_2 = os.path.join(WORKSPACE_DIR, OUTPUT_CSV_FILENAME)


def generate_october_weights():
    """
    Generate daily split weights for October 2026 (31 days).
    Standard weight structure consistent with Dynamic Split:
      - Weekday (Mon - Thu): Weight 3
      - Weekend (Fri - Sun): Weight 5
    """
    dates = pd.date_range("2026-10-01", "2026-10-31")
    weights = []
    for d in dates:
        # d.weekday(): 0=Mon, 1=Tue, 2=Wed, 3=Thu, 4=Fri, 5=Sat, 6=Sun
        w = 5.0 if d.weekday() in [4, 5, 6] else 3.0
        weights.append({
            "date": d.strftime("%Y-%m-%d"),
            "day": d.strftime("%a").lower(),
            "weight": w,
        })

    df_weights = pd.DataFrame(weights)
    total_weight = df_weights["weight"].sum()
    df_weights["ratio"] = df_weights["weight"] / total_weight
    df_weights["pct_target"] = df_weights["ratio"]
    df_weights["mtd_pct"] = df_weights["pct_target"].cumsum()

    print(f"Generated dynamic weights for October 2026: {len(df_weights)} days. Total weight: {total_weight:.0f}")
    return df_weights


def distribute_target(monthly_val, ratios, n_days):
    """
    Distribute monthly target across days using ratios.
    Adjusts the final day with remainder to guarantee 0.00 exact precision.
    """
    if monthly_val is None or pd.isna(monthly_val) or monthly_val == 0:
        return [None] * n_days

    daily_vals = []
    accumulated = 0.0
    for i in range(n_days - 1):
        val = round(monthly_val * ratios[i], 2)
        daily_vals.append(val)
        accumulated += val

    # Last day gets remainder to ensure exact match
    last_val = round(monthly_val - accumulated, 2)
    daily_vals.append(last_val)
    return daily_vals


def main():
    print("=" * 80)
    print("GENERATING SEPARATE DYNAMIC OCTOBER 2026 DAILY TARGETS (T1, T2, T3)")
    print(f"Monthly Target Source  : {MONTHLY_TARGET_FILE}")
    print("=" * 80)

    # 1. Generate October 2026 weights (31 days)
    df_weights = generate_october_weights()
    ratios = df_weights["ratio"].tolist()
    dates = df_weights["date"].tolist()
    days = df_weights["day"].tolist()
    weights_list = df_weights["weight"].tolist()
    N = len(dates)

    # 2. Load monthly targets
    df_monthly = pd.read_csv(MONTHLY_TARGET_FILE)
    print(f"Loaded {len(df_monthly)} monthly target rows.")

    target_cols = ["target_value1", "target_value2", "target_value3", "target_value4", "target_value5"]
    for col in target_cols:
        if col in df_monthly.columns:
            df_monthly[col] = pd.to_numeric(df_monthly[col], errors="coerce")
        else:
            df_monthly[col] = None

    t1_sum = df_monthly["target_value1"].dropna().sum()
    t2_sum = df_monthly["target_value2"].dropna().sum()
    t3_sum = df_monthly["target_value3"].dropna().sum()
    print(f"Monthly Target 1 Total : {t1_sum:,.2f}")
    print(f"Monthly Target 2 Total : {t2_sum:,.2f}")
    print(f"Monthly Target 3 Total : {t3_sum:,.2f}")
    print(f"Combined Monthly Total : {t1_sum + t2_sum + t3_sum:,.2f}")

    # 3. Process Daily Rows Keeping Targets Separate
    current_version = int(time.time())
    current_updated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    daily_rows_clickhouse = []
    daily_rows_summary = []
    verification_records = []

    for idx, m_row in df_monthly.iterrows():
        b_name = m_row["branch_name"]
        s_group = m_row["series_group"]

        m_t1 = m_row["target_value1"] if pd.notna(m_row["target_value1"]) and m_row["target_value1"] != 0 else None
        m_t2 = m_row["target_value2"] if pd.notna(m_row["target_value2"]) and m_row["target_value2"] != 0 else None
        m_t3 = m_row["target_value3"] if pd.notna(m_row["target_value3"]) and m_row["target_value3"] != 0 else None
        m_t4 = m_row["target_value4"] if pd.notna(m_row["target_value4"]) and m_row["target_value4"] != 0 else None
        m_t5 = m_row["target_value5"] if pd.notna(m_row["target_value5"]) and m_row["target_value5"] != 0 else None

        # Distribute each target separately
        daily_t1 = distribute_target(m_t1, ratios, N)
        daily_t2 = distribute_target(m_t2, ratios, N)
        daily_t3 = distribute_target(m_t3, ratios, N)
        daily_t4 = distribute_target(m_t4, ratios, N)
        daily_t5 = distribute_target(m_t5, ratios, N)

        # Verification calculations
        sum_d1 = sum(v for v in daily_t1 if v is not None) if m_t1 is not None else 0.0
        sum_d2 = sum(v for v in daily_t2 if v is not None) if m_t2 is not None else 0.0
        sum_d3 = sum(v for v in daily_t3 if v is not None) if m_t3 is not None else 0.0

        diff1 = round(sum_d1 - (m_t1 or 0.0), 2)
        diff2 = round(sum_d2 - (m_t2 or 0.0), 2)
        diff3 = round(sum_d3 - (m_t3 or 0.0), 2)

        is_match = (abs(diff1) < 0.001) and (abs(diff2) < 0.001) and (abs(diff3) < 0.001)

        verification_records.append({
            "branch_name": b_name,
            "series_group": s_group,
            "monthly_target_1": m_t1,
            "daily_sum_1": sum_d1,
            "diff_1": diff1,
            "monthly_target_2": m_t2,
            "daily_sum_2": sum_d2,
            "diff_2": diff2,
            "monthly_target_3": m_t3,
            "daily_sum_3": sum_d3,
            "diff_3": diff3,
            "monthly_total_all": (m_t1 or 0.0) + (m_t2 or 0.0) + (m_t3 or 0.0),
            "daily_total_all": sum_d1 + sum_d2 + sum_d3,
            "status": "EXACT MATCH" if is_match else "MISMATCH",
        })

        running_mtd1 = 0.0
        running_mtd2 = 0.0
        running_mtd3 = 0.0

        for i in range(N):
            val1 = daily_t1[i]
            val2 = daily_t2[i]
            val3 = daily_t3[i]
            val4 = daily_t4[i]
            val5 = daily_t5[i]

            running_mtd1 += (val1 or 0.0)
            running_mtd2 += (val2 or 0.0)
            running_mtd3 += (val3 or 0.0)

            # DB Schema matching ClickHouse table structure
            db_row = {
                "date": dates[i],
                "branch_name": b_name,
                "series_group": s_group,
                "target_value1": val1,
                "incentive_value1": None,
                "target_value2": val2,
                "incentive_value2": None,
                "target_value3": val3,
                "incentive_value3": None,
                "target_value4": val4,
                "incentive_value4": None,
                "target_value5": val5,
                "incentive_value5": None,
                "version": current_version,
                "updated_at": current_updated_at,
                "is_active": 1,
            }
            daily_rows_clickhouse.append(db_row)

            # Clean summary row
            total_day_val = (val1 or 0.0) + (val2 or 0.0) + (val3 or 0.0)
            total_day_mtd = running_mtd1 + running_mtd2 + running_mtd3
            summary_row = {
                "date": dates[i],
                "day": days[i],
                "day_weight": weights_list[i],
                "branch_name": b_name,
                "series_group": s_group,
                "daily_target_1": val1,
                "daily_target_2": val2,
                "daily_target_3": val3,
                "total_daily_target": total_day_val,
                "mtd_target_1": round(running_mtd1, 2),
                "mtd_target_2": round(running_mtd2, 2),
                "mtd_target_3": round(running_mtd3, 2),
                "total_mtd_target": round(total_day_mtd, 2),
            }
            daily_rows_summary.append(summary_row)

    df_db = pd.DataFrame(daily_rows_clickhouse)
    df_summary = pd.DataFrame(daily_rows_summary)
    df_verif = pd.DataFrame(verification_records)

    # 4. Verify all records
    mismatches = df_verif[df_verif["status"] != "EXACT MATCH"]
    if len(mismatches) > 0:
        print(f"CRITICAL ERROR: {len(mismatches)} branch targets did not match!")
        print(mismatches.head())
        raise ValueError("Verification failed!")
    else:
        print(f"VERIFICATION SUCCESS: All {len(df_verif)} branch records matched daily sums across Target 1, 2, 3 to 0.00 exact precision!")

    grand_total_d1 = df_db["target_value1"].sum()
    grand_total_d2 = df_db["target_value2"].sum()
    grand_total_d3 = df_db["target_value3"].sum()
    print(f"Daily Grand Total Target 1 : {grand_total_d1:,.2f} (Diff: {grand_total_d1 - t1_sum:.4f})")
    print(f"Daily Grand Total Target 2 : {grand_total_d2:,.2f} (Diff: {grand_total_d2 - t2_sum:.4f})")
    print(f"Daily Grand Total Target 3 : {grand_total_d3:,.2f} (Diff: {grand_total_d3 - t3_sum:.4f})")

    # 5. Export CSV files (for fast ClickHouse bulk insert)
    print("\nWriting CSV files...")
    for out_csv in [OUTPUT_CSV_PATH_1, OUTPUT_CSV_PATH_2]:
        os.makedirs(os.path.dirname(out_csv), exist_ok=True)
        df_db.to_csv(out_csv, index=False)
        print(f"Saved CSV: {out_csv}")

    # 6. Export to Excel with multiple clean sheets and professional styling
    print("\nWriting Excel files...")
    for out_path in [OUTPUT_EXCEL_PATH_1, OUTPUT_EXCEL_PATH_2]:
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            # Sheet 1: date_branch_wise_target (Standard ClickHouse Schema)
            df_db.to_excel(writer, sheet_name="date_branch_wise_target", index=False)

            # Sheet 2: Daily Target Summary (T1, T2, T3, MTD)
            df_summary.to_excel(writer, sheet_name="Daily Target Summary", index=False)

            # Sheet 3: Branch Monthly Summary (Verification breakdown)
            df_verif.to_excel(writer, sheet_name="Branch Monthly Summary", index=False)

            # Sheet 4: Dynamic Weights Reference
            df_weights_ref = df_weights[["date", "day", "weight", "pct_target", "mtd_pct"]].copy()
            df_weights_ref.to_excel(writer, sheet_name="Dynamic Weights Reference", index=False)

        # Post-format Excel sheets with openpyxl for clean visuals
        wb = openpyxl.load_workbook(out_path)
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws.views.sheetView[0].showGridLines = True

            # Header styling
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(1, col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Auto-adjust column widths
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = max(len(str(cell.value or "")) for cell in col[:100])
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(out_path)
        print(f"Saved Excel: {out_path}")

    print("\nGeneration completed successfully!")


if __name__ == "__main__":
    main()
